#!/usr/bin/env python3
"""Generate simple internal SVG plots for checking clinician-rating trends.

These figures are intentionally plain. They support manuscript figure planning and
are not publication-ready artwork.
"""

from __future__ import annotations

import csv
import html
from collections import Counter, defaultdict
from itertools import combinations
from pathlib import Path

from openpyxl import load_workbook


HERE = Path(__file__).resolve().parent
OUTPUT = HERE / "draft_figures"
RATER_FILES = {
    "Simo": HERE / "simo_breast_pdac_blind_scores_20260921.csv",
    "Kevin": HERE / "kevin_breast_pdac_blind_scores_20260911.csv",
    "Bolun": HERE / "bolun_breast_blind_scores_20260921.xlsx",
}
LETTER_SCORES = HERE / "patient_letter_scores_oncologist_01_summary.csv"

PL = "#2B8CBE"
BL = "#D95F0E"
TIE = "#BDBDBD"
GRID = "#E5E5E5"
TEXT = "#222222"
PURPLE = "#756BB1"

BREAST_FIELDS = [
    "current_meds",
    "stage",
    "distant_met",
    "metastasis",
    "response",
    "type_receptor",
    "genetic_results",
    "genetic_plan",
    "supportive_meds",
    "procedure_plan",
    "imaging_plan",
    "lab_plan",
    "medication_plan",
    "recent_changes",
]
PDAC_FIELDS = [field for field in BREAST_FIELDS if field != "type_receptor"]
CORE_BREAST = BREAST_FIELDS[:7]
CORE_PDAC = [
    "current_meds",
    "stage",
    "distant_met",
    "metastasis",
    "response",
    "genetic_results",
]
FIELD_LABELS = {
    "current_meds": "Active anticancer medications",
    "stage": "Stage",
    "distant_met": "Distant metastasis",
    "metastasis": "Regional / overall metastasis",
    "response": "Treatment response",
    "type_receptor": "Breast type / receptors",
    "genetic_results": "Molecular / genetic results",
}


def load_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        rows = list(csv.DictReader(handle))
    return [
        {key: (value or "").strip() for key, value in row.items()}
        for row in rows
    ]


def load_xlsx(path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["blind_scores"]
    values = sheet.iter_rows(values_only=True)
    headers = [str(value) for value in next(values)]
    rows = []
    for values_row in values:
        if not values_row[0]:
            continue
        rows.append(
            {
                key: "" if value is None else str(value).strip()
                for key, value in zip(headers, values_row)
            }
        )
    return rows


def load_raters() -> dict[str, list[dict[str, str]]]:
    return {
        name: load_xlsx(path) if path.suffix == ".xlsx" else load_csv(path)
        for name, path in RATER_FILES.items()
    }


def subset(
    rows: list[dict[str, str]], prefix: str, fields: list[str]
) -> list[dict[str, str]]:
    allowed = set(fields)
    return [
        row
        for row in rows
        if row["sample"].startswith(prefix) and row["field"] in allowed
    ]


def esc(value: object) -> str:
    return html.escape(str(value), quote=True)


def svg_start(width: int, height: int, title: str, subtitle: str = "") -> list[str]:
    out = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">',
        '<rect width="100%" height="100%" fill="white"/>',
        '<g font-family="Arial, Helvetica, sans-serif" fill="#222222">',
        f'<text x="{width / 2:.1f}" y="30" text-anchor="middle" font-size="20" font-weight="700">{esc(title)}</text>',
    ]
    if subtitle:
        out.append(
            f'<text x="{width / 2:.1f}" y="52" text-anchor="middle" font-size="12" fill="#555555">{esc(subtitle)}</text>'
        )
    return out


def svg_end(out: list[str], path: Path) -> None:
    out.extend(["</g>", "</svg>"])
    path.write_text("\n".join(out) + "\n", encoding="utf-8")


def legend(out: list[str], x: int, y: int) -> None:
    items = [(PL, "PL better"), (TIE, "Tie"), (BL, "BL better")]
    for color, label in items:
        out.append(f'<rect x="{x}" y="{y - 10}" width="16" height="12" fill="{color}"/>')
        out.append(f'<text x="{x + 22}" y="{y}" font-size="12">{esc(label)}</text>')
        x += 105


def plot_evaluator_distribution(raters: dict[str, list[dict[str, str]]]) -> None:
    groups = []
    for index, rows in enumerate(raters.values(), start=1):
        breast = subset(rows, "b", BREAST_FIELDS)
        pdac = subset(rows, "p", PDAC_FIELDS)
        if breast:
            groups.append((f"Oncologist {index}, breast", breast))
        if pdac:
            groups.append((f"Oncologist {index}, PDAC", pdac))
    width, height = 1000, 510
    left, right, top, bar_h, gap = 225, 195, 100, 52, 20
    plot_w = width - left - right
    out = svg_start(
        width,
        height,
        "Clinician preference by evaluation",
        "Required-field judgments; raw counts shown inside segments",
    )
    legend(out, left, 78)
    for tick in range(0, 101, 20):
        x = left + plot_w * tick / 100
        out.append(f'<line x1="{x:.1f}" y1="92" x2="{x:.1f}" y2="450" stroke="{GRID}"/>')
        out.append(f'<text x="{x:.1f}" y="475" text-anchor="middle" font-size="11">{tick}%</text>')
    for index, (label, rows) in enumerate(groups):
        counts = Counter(row["score"] for row in rows)
        total = len(rows)
        y = top + index * (bar_h + gap)
        out.append(f'<text x="{left - 12}" y="{y + 31}" text-anchor="end" font-size="13">{esc(label)}</text>')
        x = left
        for key, color in [("A", PL), ("TIE", TIE), ("B", BL)]:
            value = counts[key]
            segment = plot_w * value / total
            out.append(f'<rect x="{x:.1f}" y="{y}" width="{segment:.1f}" height="{bar_h}" fill="{color}"/>')
            if segment >= 34:
                out.append(
                    f'<text x="{x + segment / 2:.1f}" y="{y + 31}" text-anchor="middle" font-size="12" fill="{TEXT}">{value}</text>'
                )
            x += segment
        directional = counts["A"] + counts["B"]
        pl_share = counts["A"] / directional
        out.append(
            f'<text x="{left + plot_w + 12}" y="{y + 23}" font-size="11">'
            f'PL/BL/Tie {counts["A"]}/{counts["B"]}/{counts["TIE"]}</text>'
        )
        out.append(
            f'<text x="{left + plot_w + 12}" y="{y + 41}" font-size="10" fill="#555555">'
            f'{pl_share:.1%} of directional</text>'
        )
    out.append(f'<text x="{left + plot_w / 2:.1f}" y="502" text-anchor="middle" font-size="12">Share of required judgments</text>')
    svg_end(out, OUTPUT / "figure2_evaluator_distribution_rough.svg")


def pairwise_agreement(
    r1: list[dict[str, str]], r2: list[dict[str, str]]
) -> tuple[float, float, int]:
    m1 = {
        (row["sample"], row["field"]): row["score"]
        for row in subset(r1, "b", BREAST_FIELDS)
    }
    m2 = {
        (row["sample"], row["field"]): row["score"]
        for row in subset(r2, "b", BREAST_FIELDS)
    }
    keys = sorted(set(m1) & set(m2))
    total = len(keys)
    same = sum(m1[key] == m2[key] for key in keys)
    order = ["A", "TIE", "B"]
    counts_1 = Counter(m1[key] for key in keys)
    counts_2 = Counter(m2[key] for key in keys)
    observed = same / total
    expected = sum(counts_1[value] * counts_2[value] for value in order) / total**2
    kappa = (observed - expected) / (1 - expected)
    return observed, kappa, total


def plot_interrater_matrix(raters: dict[str, list[dict[str, str]]]) -> None:
    internal_names = list(raters)
    display_names = {
        name: f"Oncologist {index}"
        for index, name in enumerate(internal_names, start=1)
    }
    values: dict[tuple[str, str], tuple[float, float, int]] = {}
    for name_1, name_2 in combinations(internal_names, 2):
        values[(name_1, name_2)] = pairwise_agreement(raters[name_1], raters[name_2])
    width, height = 720, 600
    left, top, cell = 190, 130, 125
    out = svg_start(
        width,
        height,
        "Pairwise breast-cancer agreement",
        "Upper triangle: exact agreement; lower triangle: Cohen's kappa; n = 280 per pair",
    )
    for index, name in enumerate(internal_names):
        label = display_names[name]
        out.append(f'<text x="{left + (index + 0.5) * cell:.1f}" y="112" text-anchor="middle" font-size="13">{esc(label)}</text>')
        out.append(f'<text x="{left - 14}" y="{top + (index + 0.55) * cell:.1f}" text-anchor="end" font-size="13">{esc(label)}</text>')
    for i, row_name in enumerate(internal_names):
        for j, column_name in enumerate(internal_names):
            x, y = left + j * cell, top + i * cell
            if i == j:
                fill, label, sublabel = "#EEF3F6", "Same rater", "280 ratings"
            else:
                pair = (row_name, column_name) if (row_name, column_name) in values else (column_name, row_name)
                agreement, kappa, _ = values[pair]
                value = agreement if i < j else kappa
                intensity = min(max((value - 0.45) / 0.55, 0), 1)
                shade = int(245 - 115 * intensity)
                fill = f"rgb({shade},{shade + 12},255)"
                label = f"{agreement:.1%}" if i < j else f"κ = {kappa:.3f}"
                sublabel = "exact agreement" if i < j else "chance-adjusted"
            out.append(f'<rect x="{x}" y="{y}" width="{cell}" height="{cell}" fill="{fill}" stroke="white" stroke-width="3"/>')
            out.append(f'<text x="{x + cell / 2:.1f}" y="{y + cell / 2:.1f}" text-anchor="middle" font-size="20" font-weight="700">{esc(label)}</text>')
            out.append(f'<text x="{x + cell / 2:.1f}" y="{y + cell / 2 + 22:.1f}" text-anchor="middle" font-size="10" fill="#555555">{esc(sublabel)}</text>')
    svg_end(out, OUTPUT / "figure3_interrater_matrix_rough.svg")


def completed_rows(
    raters: dict[str, list[dict[str, str]]],
    breast_fields: list[str],
    pdac_fields: list[str],
) -> list[dict[str, str]]:
    rows = []
    for rater_rows in raters.values():
        rows.extend(subset(rater_rows, "b", breast_fields))
        rows.extend(subset(rater_rows, "p", pdac_fields))
    return rows


def plot_core_fields(raters: dict[str, list[dict[str, str]]]) -> None:
    rows = completed_rows(raters, CORE_BREAST, CORE_PDAC)
    by_field: dict[str, Counter[str]] = defaultdict(Counter)
    for row in rows:
        by_field[row["field"]][row["score"]] += 1
    fields = CORE_BREAST
    width, height = 1120, 590
    left, right, top, bar_h, gap = 275, 190, 105, 42, 20
    plot_w = width - left - right
    out = svg_start(
        width,
        height,
        "Clinician preference by core clinical category",
        "All completed clinician evaluations; percentages use applicable ratings per category",
    )
    legend(out, left, 80)
    for tick in range(0, 101, 20):
        x = left + plot_w * tick / 100
        out.append(f'<line x1="{x:.1f}" y1="95" x2="{x:.1f}" y2="535" stroke="{GRID}"/>')
        out.append(f'<text x="{x:.1f}" y="562" text-anchor="middle" font-size="11">{tick}%</text>')
    for index, field in enumerate(fields):
        counts = by_field[field]
        total = sum(counts.values())
        y = top + index * (bar_h + gap)
        out.append(f'<text x="{left - 12}" y="{y + 27}" text-anchor="end" font-size="13">{esc(FIELD_LABELS[field])}</text>')
        x = left
        for key, color in [("A", PL), ("TIE", TIE), ("B", BL)]:
            value = counts[key]
            segment = plot_w * value / total
            out.append(f'<rect x="{x:.1f}" y="{y}" width="{segment:.1f}" height="{bar_h}" fill="{color}"/>')
            x += segment
        out.append(
            f'<text x="{left + plot_w + 12}" y="{y + 27}" font-size="12">PL/BL/Tie {counts["A"]}/{counts["B"]}/{counts["TIE"]}</text>'
        )
    out.append(f'<text x="{left + plot_w / 2:.1f}" y="585" text-anchor="middle" font-size="12">Share of applicable judgments</text>')
    svg_end(out, OUTPUT / "figure4_core_fields_rough.svg")


def note_margins(
    rows: list[dict[str, str]], prefix: str, fields: list[str]
) -> dict[str, tuple[int, int]]:
    net: dict[str, int] = defaultdict(int)
    totals: dict[str, int] = defaultdict(int)
    for row in subset(rows, prefix, fields):
        sample = row["sample"]
        net[sample] += {"A": 1, "B": -1, "TIE": 0}[row["score"]]
        totals[sample] += 1
    return {
        sample: (net[sample], totals[sample])
        for sample in sorted(net, key=lambda value: int(value[1:]))
    }


def draw_margin_panel(
    out: list[str],
    x0: int,
    y0: int,
    width: int,
    height: int,
    title: str,
    margins: dict[str, tuple[int, int]],
) -> None:
    labels = list(margins)
    values = [margins[label][0] / margins[label][1] for label in labels]
    ymin, ymax = -0.2, 0.8
    left, right, top, bottom = 45, 15, 36, 55
    plot_w, plot_h = width - left - right, height - top - bottom
    out.append(f'<text x="{x0 + width / 2:.1f}" y="{y0 + 18}" text-anchor="middle" font-size="14" font-weight="700">{esc(title)}</text>')
    for tick in [-0.2, 0.0, 0.2, 0.4, 0.6, 0.8]:
        y = y0 + top + plot_h * (ymax - tick) / (ymax - ymin)
        stroke = "#888888" if tick == 0 else GRID
        out.append(f'<line x1="{x0 + left}" y1="{y:.1f}" x2="{x0 + left + plot_w}" y2="{y:.1f}" stroke="{stroke}"/>')
        out.append(f'<text x="{x0 + left - 7}" y="{y + 4:.1f}" text-anchor="end" font-size="10">{tick:.0%}</text>')
    bar_w = plot_w / len(values) * 0.72
    for index, (label, value) in enumerate(zip(labels, values)):
        x = x0 + left + (index + 0.5) * plot_w / len(values) - bar_w / 2
        zero_y = y0 + top + plot_h * ymax / (ymax - ymin)
        value_y = y0 + top + plot_h * (ymax - value) / (ymax - ymin)
        bar_y = min(zero_y, value_y)
        bar_h = abs(zero_y - value_y)
        fill = PL if value > 0 else TIE if value == 0 else BL
        out.append(f'<rect x="{x:.1f}" y="{bar_y:.1f}" width="{bar_w:.1f}" height="{max(bar_h, 2):.1f}" fill="{fill}"/>')
        out.append(f'<text x="{x + bar_w / 2:.1f}" y="{y0 + top + plot_h + 16}" text-anchor="middle" font-size="8">{esc(label)}</text>')
    out.append(f'<text x="{x0 + 12}" y="{y0 + top + plot_h / 2:.1f}" text-anchor="middle" font-size="11" transform="rotate(-90 {x0 + 12} {y0 + top + plot_h / 2:.1f})">Normalized PL-BL margin</text>')


def plot_note_margins(raters: dict[str, list[dict[str, str]]]) -> None:
    breast_rows = [row for rows in raters.values() for row in rows]
    pdac_rows = [row for rows in raters.values() for row in rows]
    breast = note_margins(breast_rows, "b", BREAST_FIELDS)
    pdac = note_margins(pdac_rows, "p", PDAC_FIELDS)
    width, height = 1260, 500
    out = svg_start(
        width,
        height,
        "Per-note clinician preference margins",
        "(PL better - BL better) / completed ratings; breast pools 3 clinicians and PDAC pools 2",
    )
    draw_margin_panel(out, 20, 72, 600, 400, "Breast cancer", breast)
    draw_margin_panel(out, 640, 72, 600, 400, "PDAC", pdac)
    svg_end(out, OUTPUT / "figure5_note_margins_rough.svg")


def plot_technical_vs_clinician(raters: dict[str, list[dict[str, str]]]) -> None:
    technical = {
        "current_meds": (8, 0, 40),
        "stage": (6, 8, 40),
        "distant_met": (11, 3, 40),
        "metastasis": (14, 4, 40),
        "response": (13, 6, 40),
        "type_receptor": (8, 5, 20),
        "genetic_results": (6, 2, 40),
    }
    clinician_rows = completed_rows(raters, CORE_BREAST, CORE_PDAC)
    by_field: dict[str, Counter[str]] = defaultdict(Counter)
    for row in clinician_rows:
        by_field[row["field"]][row["score"]] += 1
    fields = CORE_BREAST
    width, height = 1030, 570
    left, right, top, row_gap = 290, 80, 115, 57
    plot_w = width - left - right
    xmin, xmax = -0.1, 0.9
    out = svg_start(
        width,
        height,
        "Technical audit and clinician net preference rates",
        "Descriptive comparison only; reviewers and evaluated pipeline versions differ",
    )
    for tick in [-0.1, 0.0, 0.2, 0.4, 0.6, 0.8]:
        x = left + plot_w * (tick - xmin) / (xmax - xmin)
        stroke = "#888888" if tick == 0 else GRID
        out.append(f'<line x1="{x:.1f}" y1="90" x2="{x:.1f}" y2="515" stroke="{stroke}"/>')
        out.append(f'<text x="{x:.1f}" y="540" text-anchor="middle" font-size="11">{tick:.0%}</text>')
    out.append(f'<circle cx="{left + 10}" cy="73" r="6" fill="{PURPLE}"/><text x="{left + 22}" y="77" font-size="12">Technical audit</text>')
    out.append(f'<circle cx="{left + 145}" cy="73" r="6" fill="{PL}"/><text x="{left + 157}" y="77" font-size="12">Clinician ratings</text>')
    for index, field in enumerate(fields):
        y = top + index * row_gap
        ta, tb, tn = technical[field]
        tech_rate = (ta - tb) / tn
        cc = by_field[field]
        clinician_rate = (cc["A"] - cc["B"]) / sum(cc.values())
        x1 = left + plot_w * (tech_rate - xmin) / (xmax - xmin)
        x2 = left + plot_w * (clinician_rate - xmin) / (xmax - xmin)
        out.append(f'<text x="{left - 12}" y="{y + 4}" text-anchor="end" font-size="13">{esc(FIELD_LABELS[field])}</text>')
        out.append(f'<line x1="{x1:.1f}" y1="{y}" x2="{x2:.1f}" y2="{y}" stroke="#999999" stroke-width="2"/>')
        out.append(f'<circle cx="{x1:.1f}" cy="{y}" r="7" fill="{PURPLE}"/>')
        out.append(f'<circle cx="{x2:.1f}" cy="{y}" r="7" fill="{PL}"/>')
    out.append(f'<text x="{left + plot_w / 2:.1f}" y="566" text-anchor="middle" font-size="12">Net preference rate: (PL better - BL better) / applicable judgments</text>')
    svg_end(out, OUTPUT / "supplementary_figure1_technical_vs_clinician_rough.svg")


def draw_letter_difference_panel(
    out: list[str],
    x0: int,
    y0: int,
    width: int,
    height: int,
    title: str,
    differences: list[float],
) -> None:
    limit = max(1.5, max(abs(value) for value in differences) + 0.25)
    left, right, top, bottom = 55, 20, 58, 60
    plot_w, plot_h = width - left - right, height - top - bottom
    zero_y = y0 + top + plot_h / 2
    wins = sum(value > 0 for value in differences)
    ties = sum(value == 0 for value in differences)
    losses = sum(value < 0 for value in differences)
    mean_difference = sum(differences) / len(differences)
    out.append(f'<text x="{x0 + width / 2:.1f}" y="{y0 + 20}" text-anchor="middle" font-size="14" font-weight="700">{esc(title)}</text>')
    out.append(f'<text x="{x0 + width / 2:.1f}" y="{y0 + 40}" text-anchor="middle" font-size="11">Wins/ties/losses {wins}/{ties}/{losses}; mean difference {mean_difference:+.2f}</text>')
    for tick in [-limit, -limit / 2, 0, limit / 2, limit]:
        y = y0 + top + plot_h * (limit - tick) / (2 * limit)
        stroke = "#777777" if tick == 0 else GRID
        out.append(f'<line x1="{x0 + left}" y1="{y:.1f}" x2="{x0 + left + plot_w}" y2="{y:.1f}" stroke="{stroke}"/>')
        out.append(f'<text x="{x0 + left - 8}" y="{y + 4:.1f}" text-anchor="end" font-size="10">{tick:.1f}</text>')
    bar_w = plot_w / len(differences) * 0.68
    for index, value in enumerate(differences):
        x = x0 + left + (index + 0.5) * plot_w / len(differences) - bar_w / 2
        value_y = y0 + top + plot_h * (limit - value) / (2 * limit)
        fill = PL if value > 0 else TIE if value == 0 else BL
        out.append(f'<rect x="{x:.1f}" y="{min(zero_y, value_y):.1f}" width="{bar_w:.1f}" height="{max(abs(zero_y - value_y), 2):.1f}" fill="{fill}"/>')
        out.append(f'<text x="{x + bar_w / 2:.1f}" y="{y0 + top + plot_h + 16}" text-anchor="middle" font-size="8">b{index + 1}</text>')
    out.append(f'<text x="{x0 + 14}" y="{y0 + top + plot_h / 2:.1f}" text-anchor="middle" font-size="11" transform="rotate(-90 {x0 + 14} {y0 + top + plot_h / 2:.1f})">Harness minus comparator</text>')


def plot_letter_differences() -> None:
    scores = {int(row["note"]): row for row in load_csv(LETTER_SCORES)}
    if sorted(scores) != list(range(1, 21)):
        raise ValueError("Expected 20 complete breast-cancer letter-rating summaries")
    harness_vs_chatgpt = [
        float(scores[note]["harness_mean"]) - float(scores[note]["chatgpt_mean"])
        for note in range(1, 21)
    ]
    harness_vs_qwen = [
        float(scores[note]["harness_mean"]) - float(scores[note]["qwen_baseline_mean"])
        for note in range(1, 21)
    ]
    width, height = 1260, 510
    out = svg_start(
        width,
        height,
        "Exploratory patient-letter score differences",
        "Mean of accuracy, completeness, comprehensibility, and usefulness across one oncologist",
    )
    draw_letter_difference_panel(out, 20, 66, 600, 420, "Harness-based letter vs ChatGPT", harness_vs_chatgpt)
    draw_letter_difference_panel(out, 640, 66, 600, 420, "Harness-based letter vs Qwen baseline", harness_vs_qwen)
    svg_end(out, OUTPUT / "figure7_letter_differences_rough.svg")


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    raters = load_raters()
    plot_evaluator_distribution(raters)
    plot_interrater_matrix(raters)
    plot_core_fields(raters)
    plot_note_margins(raters)
    plot_technical_vs_clinician(raters)
    plot_letter_differences()
    print(f"Wrote 6 rough SVG figures to {OUTPUT}")


if __name__ == "__main__":
    main()
